"""Generate the internship submission report in Microsoft Word format."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
METRICS_PATH = BASE_DIR / "artifacts" / "test_metrics.json"
OUTPUT_PATH = BASE_DIR / "week4_report.docx"
APP_URL = "https://week-4-chhg.onrender.com/"
REPO_URL = "https://github.com/shreyas23dev/Week-4"


def main() -> None:
    metrics = json.loads(METRICS_PATH.read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(45, 55, 72)
    normal.paragraph_format.space_after = Pt(7)

    title = document.add_heading("Week 4 Capstone Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(24, 49, 83)
    subtitle = document.add_paragraph("Titanic Survival Prediction API")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(15)

    document.add_paragraph("AI Pioneers Internship — Machine Learning Internship")

    document.add_heading("Project objective", level=1)
    document.add_paragraph(
        "The goal of this capstone is to build an end-to-end supervised machine-learning "
        "application. It covers data preparation, model training and testing, model "
        "serialization with Joblib, and deployment of a prediction API using FastAPI. "
        "A browser form is included so a user can enter passenger information and see a "
        "prediction without writing an API client."
    )

    document.add_heading("Dataset and features", level=1)
    document.add_paragraph(
        "The project uses the Titanic passenger dataset, containing 891 records. The model "
        "uses passenger class, sex, age, number of siblings or spouses aboard, number of "
        "parents or children aboard, ticket fare, and embarkation port as input features. "
        "The target is the recorded survival outcome. Columns that directly repeat the "
        "target were not used as predictors."
    )

    document.add_heading("Preprocessing and model training", level=1)
    document.add_paragraph(
        "A Scikit-learn Pipeline combines preprocessing and a Random Forest classifier. "
        "Missing numeric values are imputed with training-set medians, and numeric features "
        "are standardized. Missing categorical values are filled with the most frequent "
        "training-set category, then categorical values are one-hot encoded. A stratified "
        "80/20 train/test split with a fixed random seed is used. Because transformations "
        "are inside the pipeline, their parameters are learned from training records only "
        "and then applied to the held-out test records. The complete fitted pipeline is "
        "serialized to a Joblib artifact for use by the API."
    )

    document.add_heading("Testing and evaluation", level=1)
    document.add_paragraph(
        f"The model was evaluated on {metrics['test_rows']} held-out passengers. "
        f"It achieved accuracy {metrics['accuracy']:.3f}, precision {metrics['precision']:.3f}, "
        f"recall {metrics['recall']:.3f}, F1 score {metrics['f1']:.3f}, and ROC AUC "
        f"{metrics['roc_auc']:.3f}. The saved evaluation JSON also contains the confusion "
        "matrix and per-class precision, recall, and F1. Five automated API tests pass, "
        "covering the homepage, health response, valid predictions, invalid input handling, "
        "and the response when the model artifact is unavailable. A live sample prediction "
        "was also verified through the deployed service."
    )

    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Held-out result"
    for label, key in [
        ("Accuracy", "accuracy"),
        ("Precision", "precision"),
        ("Recall", "recall"),
        ("F1 score", "f1"),
        ("ROC AUC", "roc_auc"),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = f"{metrics[key]:.3f}"

    document.add_heading("API deployment", level=1)
    document.add_paragraph(
        "The FastAPI service provides GET /health to check that the saved model artifact "
        "is available and POST /predict to accept validated passenger features and return "
        "the predicted class and probabilities. The root page hosts an interactive form, "
        "and /docs exposes the automatically generated API documentation. The service is "
        "deployed on Render and was verified at the public application URL."
    )
    document.add_paragraph(f"Live application: {APP_URL}")
    document.add_paragraph(f"API documentation: {APP_URL}docs")

    document.add_heading("Conclusion and limitations", level=1)
    document.add_paragraph(
        "This project demonstrates the full workflow from dataset preprocessing through "
        "model evaluation, serialization, and API serving. The reported scores describe "
        "one fixed test split and are not guarantees of future performance. The dataset is "
        "historical and limited; predictions should be treated as educational examples, "
        "not causal conclusions or advice about individuals. The application is not suitable "
        "for consequential decisions. Further improvements could include repeated or "
        "nested cross-validation, monitoring, model versioning, and a fairness assessment "
        "before any wider use."
    )
    document.add_paragraph(f"Source dataset: https://github.com/mwaskom/seaborn-data/blob/master/titanic.csv")
    document.add_paragraph(f"Source code: {REPO_URL}")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AI Pioneers Internship · Week 4 Capstone")

    document.save(OUTPUT_PATH)
    print(f"Saved Word report: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
